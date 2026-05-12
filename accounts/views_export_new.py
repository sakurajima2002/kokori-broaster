from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from django.template.loader import render_to_string
from django.template.loader import get_template
from xhtml2pdf import pisa
from io import BytesIO

class ExportDashboardExcelView(LoginRequiredMixin, StaffPermissionRequiredMixin, View):
    permission_required = 'accounts.view_user'

    def get(self, request, *args, **kwargs):
        start_date_str = request.GET.get('start_date')
        end_date_str = request.GET.get('end_date')
        sections = request.GET.getlist('sections')
        
        from django.utils.dateparse import parse_date
        
        wb = Workbook()
        wb.remove(wb.active) # Remove default sheet
        
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="ea580c", end_color="ea580c", fill_type="solid")
        
        # Base Orders Query
        orders_qs = Order.objects.all().prefetch_related('details__product', 'user')
        if start_date_str:
            start_date = parse_date(start_date_str)
            if start_date:
                orders_qs = orders_qs.filter(order_date__date__gte=start_date)
        if end_date_str:
            end_date = parse_date(end_date_str)
            if end_date:
                orders_qs = orders_qs.filter(order_date__date__lte=end_date)

        # 1. Resumen
        if 'summary' in sections:
            ws = wb.create_sheet(title="Resumen Financiero")
            ws.merge_cells('A1:B1')
            cell = ws['A1']
            cell.value = "Resumen de Rendimiento - Kokori Broaster"
            cell.font = Font(bold=True, size=14)
            
            ws['A3'] = "Fecha Inicio"
            ws['B3'] = start_date_str or "Todas"
            ws['A4'] = "Fecha Fin"
            ws['B4'] = end_date_str or "Todas"
            
            ws['A6'] = "Métrica"
            ws['B6'] = "Valor"
            for row in ws['A6:B6']:
                for c in row:
                    c.font = header_font
                    c.fill = header_fill
                    
            total_sales = orders_qs.exclude(status__in=['pending', 'cancelled']).aggregate(Sum('total'))['total__sum'] or 0
            total_orders = orders_qs.count()
            completed_orders = orders_qs.filter(status='delivered').count()
            
            ws['A7'] = "Ventas Totales ($)"
            ws['B7'] = float(total_sales)
            ws['A8'] = "Total de Pedidos"
            ws['B8'] = total_orders
            ws['A9'] = "Pedidos Completados"
            ws['B9'] = completed_orders
            
            ws.column_dimensions['A'].width = 25
            ws.column_dimensions['B'].width = 20

        # 2. Pedidos
        if 'orders' in sections:
            ws = wb.create_sheet(title="Detalle de Pedidos")
            headers = ['ID Pedido', 'Fecha', 'Cliente', 'Estado', 'Total ($)', 'Productos']
            ws.append(headers)
            for cell in ws[1]:
                cell.font = header_font
                cell.fill = header_fill

            for order in orders_qs.order_by('-order_date'):
                products_str = ", ".join([f"{d.quantity}x {d.product.name}" for d in order.details.all()])
                ws.append([
                    order.id,
                    order.order_date.strftime('%Y-%m-%d %H:%M'),
                    order.user.get_full_name() or order.user.username,
                    order.get_status_display(),
                    float(order.total),
                    products_str
                ])
                
            ws.column_dimensions['A'].width = 12
            ws.column_dimensions['B'].width = 20
            ws.column_dimensions['C'].width = 25
            ws.column_dimensions['D'].width = 15
            ws.column_dimensions['E'].width = 12
            ws.column_dimensions['F'].width = 50

        # 3. Inventario
        if 'stock' in sections:
            ws = wb.create_sheet(title="Inventario Actual")
            headers = ['ID Producto', 'Nombre', 'Categoría', 'Stock Disponible', 'Estado']
            ws.append(headers)
            for cell in ws[1]:
                cell.font = header_font
                cell.fill = header_fill
                
            from products.models import Product
            for product in Product.objects.all().order_by('stock'):
                estado = "Crítico" if product.stock <= 5 else "Bajo" if product.stock <= 10 else "Óptimo"
                ws.append([
                    product.id,
                    product.name,
                    product.category.name if product.category else "N/A",
                    product.stock,
                    estado
                ])
                
            ws.column_dimensions['A'].width = 12
            ws.column_dimensions['B'].width = 30
            ws.column_dimensions['C'].width = 20
            ws.column_dimensions['D'].width = 15
            ws.column_dimensions['E'].width = 15

        # 4. Calificaciones
        if 'ratings' in sections:
            ws = wb.create_sheet(title="Calificaciones Recientes")
            headers = ['ID Pedido', 'Cliente', 'Puntuación', 'Comentario', 'Fecha']
            ws.append(headers)
            for cell in ws[1]:
                cell.font = header_font
                cell.fill = header_fill
                
            from orders.models import Rating
            ratings = Rating.objects.filter(order__in=orders_qs).order_by('-rating_date')
            for r in ratings:
                ws.append([
                    r.order.id,
                    r.user.get_full_name() or r.user.username,
                    r.score,
                    r.comment or "Sin comentario",
                    r.rating_date.strftime('%Y-%m-%d')
                ])
                
            ws.column_dimensions['A'].width = 12
            ws.column_dimensions['B'].width = 25
            ws.column_dimensions['C'].width = 12
            ws.column_dimensions['D'].width = 50
            ws.column_dimensions['E'].width = 15

        if len(wb.sheetnames) == 0:
            ws = wb.create_sheet(title="Vacío")
            ws['A1'] = "No se seleccionó ninguna sección."

        response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response['Content-Disposition'] = 'attachment; filename="reporte_avanzado_dashboard.xlsx"'
        wb.save(response)
        return response


class ExportDashboardPDFView(LoginRequiredMixin, StaffPermissionRequiredMixin, View):
    permission_required = 'accounts.view_user'

    def get(self, request, *args, **kwargs):
        start_date_str = request.GET.get('start_date')
        end_date_str = request.GET.get('end_date')
        sections = request.GET.getlist('sections')
        
        from django.utils.dateparse import parse_date
        
        orders_qs = Order.objects.all().prefetch_related('details__product', 'user')
        if start_date_str:
            start_date = parse_date(start_date_str)
            if start_date:
                orders_qs = orders_qs.filter(order_date__date__gte=start_date)
        if end_date_str:
            end_date = parse_date(end_date_str)
            if end_date:
                orders_qs = orders_qs.filter(order_date__date__lte=end_date)
                
        context = {
            'start_date': start_date_str,
            'end_date': end_date_str,
            'sections': sections,
        }

        if 'summary' in sections:
            context['total_sales'] = orders_qs.exclude(status__in=['pending', 'cancelled']).aggregate(Sum('total'))['total__sum'] or 0
            context['total_orders'] = orders_qs.count()
            context['completed_orders'] = orders_qs.filter(status='delivered').count()
            context['cancelled_orders'] = orders_qs.filter(status='cancelled').count()
            
        if 'orders' in sections:
            context['orders'] = orders_qs.order_by('-order_date')
            
        if 'stock' in sections:
            from products.models import Product
            context['products'] = Product.objects.all().order_by('stock')
            
        if 'ratings' in sections:
            from orders.models import Rating
            context['ratings'] = Rating.objects.filter(order__in=orders_qs).order_by('-rating_date')

        template = get_template('staff/reports/dashboard_pdf.html')
        html = template.render(context)
        
        result = BytesIO()
        pdf = pisa.pisaDocument(BytesIO(html.encode("UTF-8")), result)
        
        if not pdf.err:
            response = HttpResponse(result.getvalue(), content_type='application/pdf')
            response['Content-Disposition'] = 'inline; filename="reporte_avanzado_dashboard.pdf"'
            return response
        return HttpResponse('Error generating PDF')
