from roles.mixins import StaffListingMixin, StaffHeaderMixin, StaffPermissionRequiredMixin
from django.views.generic import ListView, DetailView
from django.views import View
from django.contrib import messages
from django.shortcuts import render, redirect
from django.contrib.auth import login, logout
from django.contrib.auth.mixins import LoginRequiredMixin
from django.urls import reverse_lazy

from .models import User, Address
from .forms import LoginUserForm, RegisterUserForm, UserProfileForm
from products.models import Product
from orders.models import Order
from django.db.models import Sum

class LoginUserView(View):
    template_login = 'users/accounts/login.html'
    
    def get(self, request, *args, **kwargs):
        return render(request, self.template_login, {'form': LoginUserForm})
    
    def post(self, request, *args, **kwargs):
        form = LoginUserForm(request, request.POST)
        if form.is_valid():
            user = form.get_user()
            login(request, user)
            return redirect('home')
        messages.warning(request, 'El usuario o la contraseña no son válidos!')
        return render(request, self.template_login, context={'form': form})

class RegisterUserView(View):
    template_register = 'users/accounts/register.html'
    
    def get(self, request, *args, **kwargs):
        return render(request, self.template_register, {'form': RegisterUserForm})
    
    def post(self, request, *args, **kwargs):
        form = RegisterUserForm(request.POST)
        if form.is_valid():
            user = form.save()
            login(request, user)
            return redirect('home')
        messages.warning(request, 'El usuario o la contraseña no son válidos!')
        return render(request, self.template_register, context={'form': form})

class LogoutUserView(View):
    def get(self, request, *args, **kwargs):
        logout(request)
        return redirect('accounts:login')

class HomeView(StaffHeaderMixin, View):
    header_title = "Panel de Control"
    header_subtitle = "Visión General y Métricas de Rendimiento"

    def get(self, request, *args, **kwargs):
        from django.db.models import Avg
        from orders.models import Rating
        
        total_ratings = Rating.objects.count()
        avg_rating = Rating.objects.aggregate(Avg('score'))['score__avg'] or 5.0

        if request.user.is_authenticated and request.user.is_staff:
            # Metrics
            try:
                threshold = int(request.GET.get('stock_threshold', 5))
                if threshold < 0: threshold = 0
            except (ValueError, TypeError):
                threshold = 5
                
            total_users = User.objects.count()
            total_products = Product.objects.count()
            low_stock_products = Product.objects.filter(stock__lte=threshold).order_by('stock')
            low_stock_count = low_stock_products.count()
            
            # Date Filtering
            from django.utils.dateparse import parse_date
            import datetime
            
            start_date_str = request.GET.get('start_date')
            end_date_str = request.GET.get('end_date')
            
            orders_qs = Order.objects.all()
            if start_date_str:
                start_date = parse_date(start_date_str)
                if start_date:
                    orders_qs = orders_qs.filter(order_date__date__gte=start_date)
            if end_date_str:
                end_date = parse_date(end_date_str)
                if end_date:
                    orders_qs = orders_qs.filter(order_date__date__lte=end_date)
            
            total_orders = orders_qs.count()
            pending_orders_count = orders_qs.exclude(status__in=['delivered', 'cancelled']).count()
            
            total_sales = orders_qs.exclude(
                status__in=['pending', 'cancelled']
            ).aggregate(Sum('total'))['total__sum'] or 0
            
            recent_orders = orders_qs.prefetch_related('details').order_by('-order_date')[:5]
            
            rating_score = request.GET.get('rating_score')
            if rating_score and rating_score.isdigit():
                recent_ratings = Rating.objects.filter(score=int(rating_score)).order_by('-rating_date')[:5]
            else:
                recent_ratings = Rating.objects.all().order_by('-rating_date')[:5]

            context = self.get_context_data()
            context.update({
                'total_users': total_users,
                'total_products': total_products,
                'low_stock_products': low_stock_products,
                'low_stock_count': low_stock_count,
                'total_orders': total_orders,
                'pending_orders_count': pending_orders_count,
                'total_sales': total_sales,
                'recent_orders': recent_orders,
                'stock_threshold': threshold,
                'total_ratings': total_ratings,
                'avg_rating': avg_rating,
                'recent_ratings': recent_ratings,
                'rating_score': rating_score,
                'start_date': start_date_str,
                'end_date': end_date_str,
            })
            
            if request.headers.get('HX-Request'):
                action = request.GET.get('action')
                if action == 'filter_stock':
                    return render(request, 'staff/partials/stock_table.html', context)
                elif action == 'filter_ratings':
                    return render(request, 'staff/partials/recent_ratings.html', context)
                elif action == 'filter_dates':
                    return render(request, 'staff/partials/sales_card.html', context)
                
            return render(request, 'staff/dashboard.html', context)
        return render(request, 'users/home.html', {
            'total_ratings': total_ratings,
            'avg_rating': avg_rating,
        })

class UserListView(LoginRequiredMixin, StaffPermissionRequiredMixin, StaffListingMixin, ListView):
    model = User
    template_name = 'staff/accounts/user_list.html'
    context_object_name = 'users'
    permission_required = 'accounts.view_user'
    header_title = "Gestión de Usuarios"
    header_subtitle = "Control de Acceso y Perfiles de Staff"
    count_label = "Usuarios Registrados"

    def get_queryset(self):
        return super().get_queryset().exclude(id=self.request.user.id)

class UserDetailView(LoginRequiredMixin, StaffPermissionRequiredMixin, StaffHeaderMixin, DetailView):
    model = User
    template_name = 'staff/accounts/user_detail.html'
    context_object_name = 'user'
    permission_required = 'accounts.view_user'
    header_title = "Perfil de Usuario"
    header_subtitle = "Detalles de Cuenta y Actividad"
    header_show_back = True
    header_back_url = reverse_lazy('accounts:user_list')

class UserStaffStatusToggleView(LoginRequiredMixin, StaffPermissionRequiredMixin, View):
    permission_required = 'accounts.change_user'

    def post(self, request, pk):
        user = User.objects.get(pk=pk)
        if user == request.user:
            messages.error(request, "No puedes cambiar tu propio estado de staff.")
            return redirect('accounts:user_list')
            
        user.is_staff = not user.is_staff
        user.save()
        
        status_text = "ahora es Staff" if user.is_staff else "ya no es Staff"
        messages.success(request, f'El usuario {user.email} {status_text}.')
        return redirect('accounts:user_list')

class AddressListView(LoginRequiredMixin, StaffPermissionRequiredMixin, ListView):
    model = Address
    template_name = 'users/accounts/address_list.html'
    context_object_name = 'addresses'
    permission_required = 'accounts.view_address'

class AddressDetailView(LoginRequiredMixin, StaffPermissionRequiredMixin, DetailView):
    model = Address
    template_name = 'users/accounts/address_detail.html'
    context_object_name = 'address'
    permission_required = 'accounts.view_address'


class MyAccountView(LoginRequiredMixin, View):
    template_name = 'users/accounts/my_account.html'

    def get(self, request):
        form = UserProfileForm(instance=request.user)
        addresses = request.user.addresses.all()
        recent_orders = Order.objects.filter(user=request.user).select_related(
            'delivery'
        ).prefetch_related('details__product').order_by('-order_date')[:3]
        return render(request, self.template_name, {
            'form': form,
            'addresses': addresses,
            'recent_orders': recent_orders,
        })

    def post(self, request):
        form = UserProfileForm(request.POST, instance=request.user)
        if form.is_valid():
            form.save()
            messages.success(request, '¡Datos actualizados correctamente!')
            return redirect('accounts:my_account')
        addresses = request.user.addresses.all()
        recent_orders = Order.objects.filter(user=request.user).order_by('-order_date')[:3]
        return render(request, self.template_name, {
            'form': form,
            'addresses': addresses,
            'recent_orders': recent_orders,
        })


class DeleteAddressView(LoginRequiredMixin, View):
    def post(self, request, pk):
        from django.shortcuts import get_object_or_404
        address = get_object_or_404(Address, pk=pk, user=request.user)
        address.delete()
        messages.success(request, 'Dirección eliminada correctamente.')
        return redirect('accounts:my_account')


class MyOrdersView(LoginRequiredMixin, View):
    template_name = 'users/accounts/my_orders.html'

    def get(self, request):
        from orders.forms import RatingForm
        orders = Order.objects.filter(user=request.user).select_related(
            'delivery', 'address'
        ).prefetch_related('details__product').order_by('-order_date')
        return render(request, self.template_name, {
            'orders': orders,
            'rating_form': RatingForm()
        })
from django.http import JsonResponse
from django.utils import timezone
from datetime import timedelta

class DashboardChartDataView(LoginRequiredMixin, StaffPermissionRequiredMixin, View):
    permission_required = 'accounts.view_user'

    def get(self, request, *args, **kwargs):
        period = request.GET.get('period', '7d')
        today = timezone.now().date()
        import datetime
        
        if period == '30d':
            dates = [(today - timedelta(days=i)) for i in range(29, -1, -1)]
            labels = [d.strftime('%d %b') for d in dates]
            orders_filter = {'order_date__date__gte': dates[0]}
        elif period == 'this_year':
            dates = list(range(1, today.month + 1))
            labels = [datetime.date(today.year, m, 1).strftime('%b') for m in dates]
            orders_filter = {'order_date__year': today.year}
        else: # 7d
            dates = [(today - timedelta(days=i)) for i in range(6, -1, -1)]
            labels = [d.strftime('%d %b') for d in dates]
            orders_filter = {'order_date__date__gte': dates[0]}

        sales_data = []
        if period == 'this_year':
            for m in dates:
                month_total = Order.objects.filter(
                    order_date__year=today.year, order_date__month=m
                ).exclude(status__in=['pending', 'cancelled']).aggregate(Sum('total'))['total__sum'] or 0
                sales_data.append(float(month_total))
        else:
            for d in dates:
                daily_total = Order.objects.filter(
                    order_date__date=d
                ).exclude(status__in=['pending', 'cancelled']).aggregate(Sum('total'))['total__sum'] or 0
                sales_data.append(float(daily_total))

        # Orders by status
        status_counts = []
        status_labels = []
        for status_val, status_name in Order.STATUS_CHOICES:
            count = Order.objects.filter(status=status_val, **orders_filter).count()
            if count > 0:
                status_labels.append(status_name)
                status_counts.append(count)

        return JsonResponse({
            'sales': {
                'labels': labels,
                'data': sales_data,
            },
            'orders_by_status': {
                'labels': status_labels,
                'data': status_counts,
            }
        })

from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from django.template.loader import render_to_string
from django.template.loader import get_template
from xhtml2pdf import pisa
from io import BytesIO

class ExportDashboardExcelView(LoginRequiredMixin, StaffPermissionRequiredMixin, View):
    permission_required = 'accounts.view_user'

    def get(self, request, *args, **kwargs):
        start_date_str = request.GET.get('start_date')
        end_date_str = request.GET.get('end_date')
        sections = request.GET.getlist('sections')
        
        from django.utils.dateparse import parse_date
        
        wb = Workbook()
        wb.remove(wb.active) # Remove default sheet
        
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="ea580c", end_color="ea580c", fill_type="solid")
        
        # Base Orders Query
        orders_qs = Order.objects.all().prefetch_related('details__product', 'user')
        if start_date_str:
            start_date = parse_date(start_date_str)
            if start_date:
                orders_qs = orders_qs.filter(order_date__date__gte=start_date)
        if end_date_str:
            end_date = parse_date(end_date_str)
            if end_date:
                orders_qs = orders_qs.filter(order_date__date__lte=end_date)

        # 1. Resumen
        if 'summary' in sections:
            ws = wb.create_sheet(title="Resumen Financiero")
            ws.merge_cells('A1:B1')
            cell = ws['A1']
            cell.value = "Resumen de Rendimiento - Kokori Broaster"
            cell.font = Font(bold=True, size=14)
            
            ws['A3'] = "Fecha Inicio"
            ws['B3'] = start_date_str or "Todas"
            ws['A4'] = "Fecha Fin"
            ws['B4'] = end_date_str or "Todas"
            
            ws['A6'] = "Métrica"
            ws['B6'] = "Valor"
            for row in ws['A6:B6']:
                for c in row:
                    c.font = header_font
                    c.fill = header_fill
                    
            total_sales = orders_qs.exclude(status__in=['pending', 'cancelled']).aggregate(Sum('total'))['total__sum'] or 0
            total_orders = orders_qs.count()
            completed_orders = orders_qs.filter(status='delivered').count()
            
            ws['A7'] = "Ventas Totales ($)"
            ws['B7'] = float(total_sales)
            ws['A8'] = "Total de Pedidos"
            ws['B8'] = total_orders
            ws['A9'] = "Pedidos Completados"
            ws['B9'] = completed_orders
            
            ws.column_dimensions['A'].width = 25
            ws.column_dimensions['B'].width = 20

        # 2. Pedidos
        if 'orders' in sections:
            ws = wb.create_sheet(title="Detalle de Pedidos")
            headers = ['ID Pedido', 'Fecha', 'Cliente', 'Estado', 'Total ($)', 'Productos']
            ws.append(headers)
            for cell in ws[1]:
                cell.font = header_font
                cell.fill = header_fill

            for order in orders_qs.order_by('-order_date'):
                products_str = ", ".join([f"{d.quantity}x {d.product.name}" for d in order.details.all()])
                ws.append([
                    order.id,
                    order.order_date.strftime('%Y-%m-%d %H:%M'),
                    order.user.get_full_name() or order.user.username,
                    order.get_status_display(),
                    float(order.total),
                    products_str
                ])
                
            ws.column_dimensions['A'].width = 12
            ws.column_dimensions['B'].width = 20
            ws.column_dimensions['C'].width = 25
            ws.column_dimensions['D'].width = 15
            ws.column_dimensions['E'].width = 12
            ws.column_dimensions['F'].width = 50

        # 3. Inventario
        if 'stock' in sections:
            ws = wb.create_sheet(title="Inventario Actual")
            headers = ['ID Producto', 'Nombre', 'Categoría', 'Stock Disponible', 'Estado']
            ws.append(headers)
            for cell in ws[1]:
                cell.font = header_font
                cell.fill = header_fill
                
            from products.models import Product
            for product in Product.objects.all().order_by('stock'):
                estado = "Crítico" if product.stock <= 5 else "Bajo" if product.stock <= 10 else "Óptimo"
                ws.append([
                    product.id,
                    product.name,
                    product.category.name if product.category else "N/A",
                    product.stock,
                    estado
                ])
                
            ws.column_dimensions['A'].width = 12
            ws.column_dimensions['B'].width = 30
            ws.column_dimensions['C'].width = 20
            ws.column_dimensions['D'].width = 15
            ws.column_dimensions['E'].width = 15

        # 4. Calificaciones
        if 'ratings' in sections:
            ws = wb.create_sheet(title="Calificaciones Recientes")
            headers = ['ID Pedido', 'Cliente', 'Puntuación', 'Comentario', 'Fecha']
            ws.append(headers)
            for cell in ws[1]:
                cell.font = header_font
                cell.fill = header_fill
                
            from orders.models import Rating
            ratings = Rating.objects.filter(order__in=orders_qs).order_by('-rating_date')
            for r in ratings:
                ws.append([
                    r.order.id,
                    r.user.get_full_name() or r.user.username,
                    r.score,
                    r.comment or "Sin comentario",
                    r.rating_date.strftime('%Y-%m-%d')
                ])
                
            ws.column_dimensions['A'].width = 12
            ws.column_dimensions['B'].width = 25
            ws.column_dimensions['C'].width = 12
            ws.column_dimensions['D'].width = 50
            ws.column_dimensions['E'].width = 15

        if len(wb.sheetnames) == 0:
            ws = wb.create_sheet(title="Vacío")
            ws['A1'] = "No se seleccionó ninguna sección."

        response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response['Content-Disposition'] = 'attachment; filename="reporte_avanzado_dashboard.xlsx"'
        wb.save(response)
        return response

class ExportDashboardPDFView(LoginRequiredMixin, StaffPermissionRequiredMixin, View):
    permission_required = 'accounts.view_user'

    def get(self, request, *args, **kwargs):
        start_date_str = request.GET.get('start_date')
        end_date_str = request.GET.get('end_date')
        sections = request.GET.getlist('sections')
        
        from django.utils.dateparse import parse_date
        
        orders_qs = Order.objects.all().prefetch_related('details__product', 'user')
        if start_date_str:
            start_date = parse_date(start_date_str)
            if start_date:
                orders_qs = orders_qs.filter(order_date__date__gte=start_date)
        if end_date_str:
            end_date = parse_date(end_date_str)
            if end_date:
                orders_qs = orders_qs.filter(order_date__date__lte=end_date)
                
        context = {
            'start_date': start_date_str,
            'end_date': end_date_str,
            'sections': sections,
        }

        if 'summary' in sections:
            context['total_sales'] = orders_qs.exclude(status__in=['pending', 'cancelled']).aggregate(Sum('total'))['total__sum'] or 0
            context['total_orders'] = orders_qs.count()
            context['completed_orders'] = orders_qs.filter(status='delivered').count()
            context['cancelled_orders'] = orders_qs.filter(status='cancelled').count()
            
        if 'orders' in sections:
            context['orders'] = orders_qs.order_by('-order_date')
            
        if 'stock' in sections:
            from products.models import Product
            context['products'] = Product.objects.all().order_by('stock')
            
        if 'ratings' in sections:
            from orders.models import Rating
            context['ratings'] = Rating.objects.filter(order__in=orders_qs).order_by('-rating_date')

        template = get_template('staff/reports/dashboard_pdf.html')
        html = template.render(context)
        
        result = BytesIO()
        pdf = pisa.pisaDocument(BytesIO(html.encode("UTF-8")), result)
        
        if not pdf.err:
            response = HttpResponse(result.getvalue(), content_type='application/pdf')
            response['Content-Disposition'] = 'inline; filename="reporte_avanzado_dashboard.pdf"'
            return response
        return HttpResponse('Error generating PDF')
